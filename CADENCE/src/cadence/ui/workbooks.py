"""Workbook parsing and provenance-preserving portfolio selection."""

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Iterable, Union

import polars as pl
from openpyxl import load_workbook

from cadence.reference_data.economics_assets import read_economics_asset_workbook
from cadence.reference_data.year1_damage import read_asset_workbook
from cadence.ui.paths import MATERIAL_CLASS_MAP, MATERIAL_MAPPING

ASSET_SHEET = "Sheet1"
MATERIAL_LABELS = {
    "OFFICIAL_ASPHALT": "Asphalt",
    "OFFICIAL_METAL": "Metal",
    "OFFICIAL_TILE": "Tile",
}


@dataclass(frozen=True)
class ValidatedPortfolio:
    """Source records plus normalized fields from authoritative CADENCE readers."""

    source: pl.DataFrame
    normalized: pl.DataFrame

    @property
    def display(self) -> pl.DataFrame:
        normalized = self.normalized.select(
            "asset_id",
            "official_current_material_id",
            "input_roof_age",
            "terrain_id",
            "terrain_label",
        ).with_columns(
            pl.col("official_current_material_id")
            .replace_strict(MATERIAL_LABELS)
            .alias("official_material")
        )
        source = self.source.with_columns(
            pl.col("asset_id").cast(pl.String).str.strip_chars()
        )
        return source.join(normalized, on="asset_id", how="left", validate="1:1")


def sha256_file(path: Path) -> str:
    """Return the SHA-256 identity of one file."""
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def save_uploaded_workbook(upload: BinaryIO, destination: Path) -> Path:
    """Atomically persist a browser upload without modifying repository inputs."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(".xlsx.tmp")
    upload.seek(0)
    with temporary.open("wb") as target:
        for chunk in iter(lambda: upload.read(1024 * 1024), b""):
            target.write(chunk)
    temporary.replace(destination)
    return destination


def validate_portfolio(path: Path, sheet_name: str = ASSET_SHEET) -> ValidatedPortfolio:
    """Validate a workbook through both authoritative CADENCE ingestion contracts."""
    economics = read_economics_asset_workbook(
        path, MATERIAL_MAPPING, MATERIAL_CLASS_MAP, sheet_name
    )
    damage = read_asset_workbook(path, sheet_name)
    source = _read_source_sheet(path, sheet_name)
    normalized = economics.select(
        "asset_id", "official_current_material_id", "input_roof_age"
    ).join(
        damage.select("asset_id", "terrain_id", "terrain_label"),
        on="asset_id",
        how="inner",
        validate="1:1",
    )
    if normalized.height != source.height:
        raise ValueError("validated asset rows do not match source workbook rows")
    return ValidatedPortfolio(source=source, normalized=normalized)


def analysis_identity(source_checksum: str, asset_ids: Iterable[str]) -> str:
    """Identify a selected portfolio independently of display filter ordering."""
    payload = json.dumps(
        {"source_checksum": source_checksum, "asset_ids": sorted(set(asset_ids))},
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def materialize_filtered_workbook(
    source_path: Path,
    destination: Path,
    asset_ids: Iterable[str],
    sheet_name: str = ASSET_SHEET,
) -> Path:
    """Copy a workbook and retain only selected asset rows in its asset sheet."""
    selected = {str(asset_id).strip() for asset_id in asset_ids}
    if not selected:
        raise ValueError("at least one asset must be selected for analysis")
    workbook = load_workbook(source_path)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"asset sheet {sheet_name!r} not found")
    sheet = workbook[sheet_name]
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    if "asset_id" not in headers:
        raise ValueError("asset workbook is missing required column: asset_id")
    asset_column = headers.index("asset_id") + 1
    present = {
        str(sheet.cell(row=row, column=asset_column).value).strip()
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=asset_column).value is not None
    }
    missing = sorted(selected - present)
    if missing:
        raise ValueError("selected asset_ids are absent from workbook: " + ", ".join(missing))
    for row in range(sheet.max_row, 1, -1):
        value = sheet.cell(row=row, column=asset_column).value
        if str(value).strip() not in selected:
            sheet.delete_rows(row)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(".xlsx.tmp")
    workbook.save(temporary)
    temporary.replace(destination)
    return destination


def _read_source_sheet(path: Path, sheet_name: str) -> pl.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"asset sheet {sheet_name!r} not found")
    rows = workbook[sheet_name].iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as error:
        raise ValueError("asset workbook is empty") from error
    records = [
        dict(zip(headers, row))
        for row in rows
        if any(value is not None and str(value).strip() for value in row)
    ]
    return pl.DataFrame(records, strict=False)