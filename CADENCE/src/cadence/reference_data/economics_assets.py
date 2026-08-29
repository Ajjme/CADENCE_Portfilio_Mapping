"""Build economics-ready asset features from the CADENCE asset workbook."""

import hashlib
import json
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, Optional, Sequence

import duckdb
import geopandas as gpd
import pandas as pd
import polars as pl
from openpyxl import load_workbook

ECONOMICS_ASSET_SCHEMA_VERSION = "v0.1.0"
MAX_NEAREST_GEOGRAPHY_DISTANCE_M = 50_000.0
GEOGRAPHY_DISTANCE_CRS = "EPSG:5070"
WORKBOOK_COLUMNS = (
    "asset_id",
    "latitude",
    "longitude",
    "current_roof_type",
    "roof_age",
    "roof_area_sqft",
    "roof_shape",
    "roof_deck_attachment",
    "roof_wall_connection",
)
GEO_OUTPUT_COLUMNS = (
    "asset_id",
    "zip_code",
    "county_fips",
    "cbsa_code",
    "state_code",
    "labor_market_id",
)
FIPS_TO_STATE = {
    "01": "AL", "02": "AK", "04": "AZ", "05": "AR", "06": "CA",
    "08": "CO", "09": "CT", "10": "DE", "11": "DC", "12": "FL",
    "13": "GA", "15": "HI", "16": "ID", "17": "IL", "18": "IN",
    "19": "IA", "20": "KS", "21": "KY", "22": "LA", "23": "ME",
    "24": "MD", "25": "MA", "26": "MI", "27": "MN", "28": "MS",
    "29": "MO", "30": "MT", "31": "NE", "32": "NV", "33": "NH",
    "34": "NJ", "35": "NM", "36": "NY", "37": "NC", "38": "ND",
    "39": "OH", "40": "OK", "41": "OR", "42": "PA", "44": "RI",
    "45": "SC", "46": "SD", "47": "TN", "48": "TX", "49": "UT",
    "50": "VT", "51": "VA", "53": "WA", "54": "WV", "55": "WI",
    "56": "WY", "72": "PR",
}
DECK_ATTACHMENT_ALIASES = {
    '6d @ 6"/12"': "6d_6in_12in",
    '8d @ 6"/12"': "8d_6in_12in",
    '8d @ 6"/6"': "8d_6in_6in",
    '6d/8d @ 6"/6"': "6d_8d_mix_6in_6in",
}


def read_economics_asset_workbook(
    workbook_path: Path,
    mapping_path: Path,
    class_map_path: Path,
    sheet_name: str = "Sheet1",
) -> pl.DataFrame:
    """Read user roof fields and resolve them to the economics asset contract."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"asset sheet {sheet_name!r} not found")
    rows = workbook[sheet_name].iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as error:
        raise ValueError("asset workbook is empty") from error
    missing = sorted(set(WORKBOOK_COLUMNS) - set(headers))
    if missing:
        raise ValueError(f"asset workbook is missing economics columns: {missing}")
    records = [
        dict(zip(headers, row))
        for row in rows
        if any(value is not None and str(value).strip() for value in row)
    ]
    if not records:
        raise ValueError("asset workbook contains no asset rows")

    normalized = []
    seen = set()
    for row_number, record in enumerate(records, start=2):
        asset_id = str(record["asset_id"]).strip()
        if not asset_id or asset_id in seen:
            raise ValueError(f"row {row_number}: asset_id must be nonblank and unique")
        seen.add(asset_id)
        latitude = _number(record["latitude"], row_number, "latitude")
        longitude = _number(record["longitude"], row_number, "longitude")
        roof_age = _number(record["roof_age"], row_number, "roof_age")
        roof_area = _number(record["roof_area_sqft"], row_number, "roof_area_sqft")
        if not -90 <= latitude <= 90 or not -180 <= longitude <= 180:
            raise ValueError(f"row {row_number}: coordinates are outside WGS84 bounds")
        if not roof_age.is_integer() or roof_age < 1:
            raise ValueError(f"row {row_number}: roof_age must be a positive integer")
        if roof_area <= 0:
            raise ValueError(f"row {row_number}: roof_area_sqft must be positive")
        normalized.append(
            {
                "asset_id": asset_id,
                "latitude": latitude,
                "longitude": longitude,
                "raw_current_roof_type": str(record["current_roof_type"]).strip(),
                "normalized_current_roof_type": str(record["current_roof_type"])
                .strip()
                .casefold(),
                "input_roof_age": int(roof_age),
                "roof_area_sqft": roof_area,
                "roof_shape": _normalize_token(record["roof_shape"]),
                "roof_deck_attachment": _normalize_deck_attachment(
                    record["roof_deck_attachment"]
                ),
                "roof_wall_connection": _normalize_token(
                    record["roof_wall_connection"]
                ).replace("-", "_"),
            }
        )
    assets = pl.DataFrame(normalized)
    material_map = _asset_material_map(mapping_path, class_map_path)
    assets = assets.join(
        material_map,
        left_on="normalized_current_roof_type",
        right_on="normalized_source_value",
        how="left",
        validate="m:1",
    )
    unresolved = assets.filter(pl.col("official_current_material_id").is_null())
    if unresolved.height:
        values = sorted(unresolved["raw_current_roof_type"].unique().to_list())
        raise ValueError("unsupported current_roof_type values: " + ", ".join(values))
    return assets.sort("asset_id")


def attach_economics_geography(
    assets: pl.DataFrame,
    zcta_path: Path,
    zcta_crosswalk_path: Path,
    cbsa_crosswalk_zip_path: Path,
    labor_geographies_path: Path,
    labor_wages_path: Path,
) -> pl.DataFrame:
    """Attach ZCTA, county, CBSA, state, and BLS labor area to asset points."""
    _require_columns(assets, ("asset_id", "latitude", "longitude"), "assets")
    points = gpd.GeoDataFrame(
        assets.select("asset_id", "latitude", "longitude").to_pandas(),
        geometry=gpd.points_from_xy(assets["longitude"], assets["latitude"]),
        crs="EPSG:4326",
    )
    bounds = tuple(points.total_bounds)
    zcta_source = f"zip://{zcta_path.resolve()}" if zcta_path.suffix == ".zip" else zcta_path
    zctas = gpd.read_file(zcta_source, bbox=bounds)[["ZCTA5CE20", "geometry"]]
    zctas = zctas.to_crs("EPSG:4326").rename(columns={"ZCTA5CE20": "zip_code"})
    point_zcta = _assign_polygon_geography(
        points,
        zctas,
        "zip_code",
        "ZCTA",
        "tiger_zcta_point_intersection",
        "tiger_zcta_nearest_within_50km",
    )

    labor = gpd.read_parquet(labor_geographies_path)[["AREA", "geometry"]]
    labor = labor.to_crs("EPSG:4326").rename(columns={"AREA": "labor_market_id"})
    point_labor = _assign_polygon_geography(
        points,
        labor,
        "labor_market_id",
        "labor market",
        "bls_area_point_intersection",
        "bls_area_nearest_within_50km",
    )
    labor_ids = pl.read_parquet(labor_wages_path).select(
        pl.col("AREA").alias("labor_market_id")
    ).unique()

    spatial = point_zcta.join(
        point_labor, on="asset_id", how="inner", validate="1:1"
    ).with_columns(
        pl.col("zip_code").cast(pl.String).str.zfill(5),
        pl.col("labor_market_id").cast(pl.String),
    )
    if spatial.join(labor_ids, on="labor_market_id", how="anti").height:
        raise ValueError("one or more assigned labor markets have no wage data")
    crosswalk = build_selected_zcta_crosswalk(
        spatial["zip_code"].unique().to_list(),
        zcta_crosswalk_path,
        cbsa_crosswalk_zip_path,
    )
    result = spatial.join(crosswalk, on="zip_code", how="left", validate="m:1")
    result = result.with_columns(
        pl.col("county_fips").str.slice(0, 2).replace_strict(FIPS_TO_STATE).alias(
            "state_code"
        ),
        pl.lit("nhgis_dominant_block_weight").alias("county_assignment_method"),
        pl.lit("nhgis_joint_block_weight").alias("cbsa_assignment_method"),
        pl.lit("v2020").alias("geography_version"),
    )
    unresolved = result.filter(
        pl.any_horizontal(
            pl.col(column).is_null()
            for column in ("county_fips", "cbsa_code", "state_code", "labor_market_id")
        )
    )
    if unresolved.height:
        raise ValueError(
            "economics geography could not be resolved for asset_ids: "
            + ", ".join(unresolved["asset_id"].to_list())
        )
    return assets.join(result, on="asset_id", how="left", validate="1:1").sort(
        "asset_id"
    )


def _assign_polygon_geography(
    points: gpd.GeoDataFrame,
    polygons: gpd.GeoDataFrame,
    value_column: str,
    geography_name: str,
    intersection_method: str,
    nearest_method: str,
    max_distance_m: float = MAX_NEAREST_GEOGRAPHY_DISTANCE_M,
) -> pl.DataFrame:
    """Assign polygons directly, then use a bounded nearest match for true gaps."""
    joined = gpd.sjoin(
        points,
        polygons[[value_column, "geometry"]],
        how="left",
        predicate="intersects",
    )
    matched = joined[joined[value_column].notna()][["asset_id", value_column]]
    if matched["asset_id"].duplicated().any():
        raise ValueError(f"an asset point matched more than one {geography_name} polygon")
    matched = matched.assign(
        assignment_method=intersection_method,
        assignment_distance_m=0.0,
    )
    unresolved = points[~points["asset_id"].isin(matched["asset_id"])]
    assignments = [matched]
    if not unresolved.empty:
        nearest = gpd.sjoin_nearest(
            unresolved.to_crs(GEOGRAPHY_DISTANCE_CRS),
            polygons.to_crs(GEOGRAPHY_DISTANCE_CRS)[[value_column, "geometry"]],
            how="left",
            distance_col="assignment_distance_m",
        )
        nearest = nearest.sort_values(
            ["asset_id", "assignment_distance_m", value_column]
        ).drop_duplicates("asset_id", keep="first")
        failed = nearest[
            nearest[value_column].isna()
            | (nearest["assignment_distance_m"] > max_distance_m)
        ]
        if not failed.empty:
            details = ", ".join(
                f"{row.asset_id} ({row.assignment_distance_m / 1000:.1f} km)"
                for row in failed.itertuples()
            )
            raise ValueError(
                f"{geography_name} could not be resolved within "
                f"{max_distance_m / 1000:.0f} km for asset_ids: {details}"
            )
        assignments.append(
            nearest[["asset_id", value_column, "assignment_distance_m"]].assign(
                assignment_method=nearest_method
            )
        )
    combined = pd.concat(assignments, ignore_index=True)
    if len(combined) != len(points) or combined["asset_id"].duplicated().any():
        raise ValueError(f"{geography_name} assignment did not preserve asset cardinality")
    prefix = "zip" if value_column == "zip_code" else "labor"
    return pl.from_pandas(combined).rename(
        {
            "assignment_method": f"{prefix}_assignment_method",
            "assignment_distance_m": f"{prefix}_assignment_distance_m",
        }
    )


def build_selected_zcta_crosswalk(
    zip_codes: Sequence[str],
    zcta_crosswalk_path: Path,
    cbsa_crosswalk_zip_path: Path,
) -> pl.DataFrame:
    """Resolve dominant county and CBSA for only the selected ZCTAs."""
    selected = sorted(set(zip_codes))
    if not selected:
        raise ValueError("at least one ZIP code is required")
    sql_values = ", ".join(f"'{_sql_literal(value)}'" for value in selected)
    with tempfile.TemporaryDirectory(prefix="cadence-cbsa-") as temporary:
        with zipfile.ZipFile(cbsa_crosswalk_zip_path) as archive:
            members = [name for name in archive.namelist() if name.endswith(".csv")]
            if len(members) != 1:
                raise ValueError("CBSA crosswalk archive must contain exactly one CSV")
            cbsa_path = Path(archive.extract(members[0], temporary))
        query = f"""
            WITH zcta AS (
                SELECT
                    blk2010ge,
                    LPAD(CAST(zcta2020ge AS VARCHAR), 5, '0') AS zip_code,
                    CAST(weight AS DOUBLE) AS zcta_weight
                FROM read_csv(
                    '{_sql_path(zcta_crosswalk_path)}',
                    header=true,
                    all_varchar=true
                )
                WHERE LPAD(CAST(zcta2020ge AS VARCHAR), 5, '0') IN ({sql_values})
            ),
            county_scores AS (
                SELECT zip_code, SUBSTR(blk2010ge, 1, 5) AS county_fips,
                       SUM(zcta_weight) AS county_crosswalk_weight
                FROM zcta
                GROUP BY zip_code, county_fips
            ),
            county AS (
                SELECT * EXCLUDE (rank)
                FROM (
                    SELECT *, ROW_NUMBER() OVER (
                        PARTITION BY zip_code
                        ORDER BY county_crosswalk_weight DESC, county_fips
                    ) AS rank
                    FROM county_scores
                )
                WHERE rank = 1
            ),
            cbsa_source AS (
                SELECT blk2010ge,
                       LPAD(CAST(cbsa2020ge AS VARCHAR), 5, '0') AS cbsa_code,
                       CAST(weight AS DOUBLE) AS cbsa_weight
                FROM read_csv(
                    '{_sql_path(cbsa_path)}',
                    header=true,
                    all_varchar=true
                )
                WHERE cbsa2020ge IS NOT NULL
            ),
            cbsa_scores AS (
                SELECT z.zip_code, c.cbsa_code,
                       SUM(z.zcta_weight * c.cbsa_weight) AS cbsa_crosswalk_weight
                FROM zcta z
                JOIN cbsa_source c USING (blk2010ge)
                GROUP BY z.zip_code, c.cbsa_code
            ),
            cbsa AS (
                SELECT * EXCLUDE (rank)
                FROM (
                    SELECT *, ROW_NUMBER() OVER (
                        PARTITION BY zip_code
                        ORDER BY cbsa_crosswalk_weight DESC, cbsa_code
                    ) AS rank
                    FROM cbsa_scores
                )
                WHERE rank = 1
            )
            SELECT county.zip_code, county.county_fips,
                   county.county_crosswalk_weight,
                   cbsa.cbsa_code, cbsa.cbsa_crosswalk_weight
            FROM county
            LEFT JOIN cbsa USING (zip_code)
            ORDER BY county.zip_code
        """
        result = duckdb.sql(query).pl()
    missing = sorted(set(selected) - set(result["zip_code"].to_list()))
    if missing:
        raise ValueError("NHGIS crosswalk is missing ZCTAs: " + ", ".join(missing))
    return result


def build_economics_asset_features(
    workbook_path: Path,
    repository_root: Path,
    output_path: Path,
    sheet_name: str = "Sheet1",
) -> Dict[str, object]:
    """Build and atomically publish one economics-ready asset Parquet."""
    paths = _source_paths(repository_root)
    assets = read_economics_asset_workbook(
        workbook_path, paths["mapping"], paths["class_map"], sheet_name
    )
    features = attach_economics_geography(
        assets,
        paths["zcta"],
        paths["zcta_crosswalk"],
        paths["cbsa_crosswalk"],
        paths["labor_geographies"],
        paths["labor_wages"],
    )
    _validate_output(features)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_output = output_path.with_suffix(".parquet.tmp")
    features.write_parquet(temporary_output)
    temporary_output.replace(output_path)
    manifest_path = output_path.with_suffix(".manifest.json")
    manifest = {
        "schema_version": ECONOMICS_ASSET_SCHEMA_VERSION,
        "build_timestamp": datetime.now(timezone.utc).isoformat(),
        "source_workbook": str(workbook_path.resolve()),
        "source_workbook_sha256": _sha256(workbook_path),
        "asset_count": features.height,
        "geography_version": "v2020",
        "zip_assignment_method_counts": _value_counts(
            features, "zip_assignment_method"
        ),
        "zip_assignment_max_distance_m": features[
            "zip_assignment_distance_m"
        ].max(),
        "labor_assignment_method_counts": _value_counts(
            features, "labor_assignment_method"
        ),
        "labor_assignment_max_distance_m": features[
            "labor_assignment_distance_m"
        ].max(),
        "output_path": str(output_path.resolve()),
    }
    temporary_manifest = manifest_path.with_suffix(".json.tmp")
    temporary_manifest.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    temporary_manifest.replace(manifest_path)
    return manifest


def _asset_material_map(mapping_path: Path, class_map_path: Path) -> pl.DataFrame:
    mapping = pl.read_csv(mapping_path).filter(
        (pl.col("mapping_domain") == "roof_material")
        & (pl.col("source_dataset_id") == "asset_inventory_test")
        & (pl.col("source_column") == "current_roof_type")
        & pl.col("mapping_status").str.starts_with("approved")
    ).select(
        "normalized_source_value",
        "canonical_id",
        (pl.col("relationship") == "explicit_default").alias("default_applied"),
        pl.col("mapping_id").alias("material_mapping_id"),
    )
    classes = pl.read_csv(class_map_path).filter(
        pl.col("mapping_version") == "v1.0.0"
    ).select(
        "canonical_id",
        pl.col("official_class_id").alias("official_current_material_id"),
        pl.col("mapping_version").alias("material_mapping_version"),
    )
    return mapping.join(classes, on="canonical_id", how="left", validate="1:1")


def _value_counts(frame: pl.DataFrame, column: str) -> Dict[str, int]:
    return {
        str(row[column]): int(row["len"])
        for row in frame.group_by(column).len().iter_rows(named=True)
    }


def _validate_output(features: pl.DataFrame) -> None:
    required = (
        "asset_id", "official_current_material_id", "input_roof_age",
        "roof_area_sqft", "zip_code", "cbsa_code",
        "state_code", "county_fips", "labor_market_id", "roof_shape",
        "roof_deck_attachment", "roof_wall_connection",
    )
    _require_columns(features, required, "economics asset features")
    if features.select(required).null_count().sum_horizontal().sum() > 0:
        raise ValueError("economics asset features contain null required values")


def _source_paths(repository_root: Path) -> Dict[str, Path]:
    geography = repository_root / "Data" / "Geography"
    mapping = repository_root / "Data_Catalogs" / "Mapping"
    labor = repository_root / "Data" / "Labor" / "Start_Year_2026"
    return {
        "zcta": geography / "Census_TIGER" / "raw" / "v2020" / "tl_2020_us_zcta520.zip",
        "zcta_crosswalk": geography / "IPUMS_NHGIS" / "raw" / "v2020" / "nhgis_blk2010_zcta2020" / "nhgis_blk2010_zcta2020.csv",
        "cbsa_crosswalk": geography / "IPUMS_NHGIS" / "raw" / "v2020" / "nhgis_blk2010_cbsa2020.zip",
        "labor_geographies": labor / "labor_geographies.parquet",
        "labor_wages": labor / "labor_wages_long.parquet",
        "mapping": mapping / "master_mapping_reference_draft.csv",
        "class_map": mapping / "official_material_class_map_v1.csv",
    }


def _normalize_token(value: object) -> str:
    return str(value).strip().casefold().replace(" ", "_")


def _normalize_deck_attachment(value: object) -> str:
    text = str(value).strip()
    return DECK_ATTACHMENT_ALIASES.get(text, _normalize_token(text))


def _number(value: object, row_number: int, field: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"row {row_number}: {field} must be numeric") from error
    if not (-float("inf") < result < float("inf")):
        raise ValueError(f"row {row_number}: {field} must be finite")
    return result


def _require_columns(frame: pl.DataFrame, columns: Iterable[str], name: str) -> None:
    missing = sorted(set(columns) - set(frame.columns))
    if missing:
        raise ValueError(f"{name} is missing required columns: {missing}")


def _sql_literal(value: str) -> str:
    return value.replace("'", "''")


def _sql_path(path: Path) -> str:
    return str(path.resolve()).replace("'", "''")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()