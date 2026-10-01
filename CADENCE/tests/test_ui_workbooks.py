from pathlib import Path

from openpyxl import load_workbook

from cadence.economics.contracts import CostStream
from cadence.ui.pages import portfolio
from cadence.ui.state import update_analysis_identity
from cadence.ui.paths import DEFAULT_WORKBOOK
from cadence.ui.workbooks import (
    analysis_identity,
    materialize_filtered_workbook,
    sha256_file,
    validate_portfolio,
)


def test_default_workbook_uses_asset_sheet_and_official_materials() -> None:
    portfolio = validate_portfolio(DEFAULT_WORKBOOK)

    assert portfolio.source.height == 6
    assert set(portfolio.display["official_material"].to_list()) <= {
        "Asphalt", "Metal", "Tile"
    }
    assert "current_roof_type" in portfolio.source.columns


def test_portfolio_download_offers_default_workbook(monkeypatch) -> None:
    download = {}
    monkeypatch.setattr(portfolio.st, "markdown", lambda *args, **kwargs: None)
    monkeypatch.setattr(portfolio.st, "info", lambda *args, **kwargs: None)
    monkeypatch.setattr(portfolio.st, "file_uploader", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        portfolio.st,
        "download_button",
        lambda label, **kwargs: download.update(label=label, **kwargs),
    )

    source, _, _ = portfolio._portfolio_source()

    assert source == DEFAULT_WORKBOOK
    assert download["data"] == DEFAULT_WORKBOOK.read_bytes()
    assert download["file_name"] == DEFAULT_WORKBOOK.name
    assert download["mime"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_demand_surge_settings_change_run_identity(monkeypatch) -> None:
    config = portfolio.load_economics_config()
    monkeypatch.setattr(portfolio.st, "number_input", lambda label, **kwargs: kwargs["value"])
    monkeypatch.setattr(
        portfolio.st, "selectbox", lambda label, options, **kwargs: options[kwargs["index"]]
    )
    monkeypatch.setattr(
        portfolio.st, "segmented_control", lambda label, options, **kwargs: kwargs["default"]
    )
    settings = portfolio._analysis_settings(config)
    assert settings["demand_surge"] is False

    identity = portfolio._effective_identity("selected-assets", config, settings)
    state = {"analysis_identity": identity, "run_state": {"run_id": "old"}}
    surged_settings = {**settings, "demand_surge": True}
    surged_identity = portfolio._effective_identity("selected-assets", config, surged_settings)
    assert surged_identity != identity
    assert update_analysis_identity(state, surged_identity) is True
    assert "run_state" not in state
    assert portfolio._configured_run(config, surged_settings).demand_surge is True


def test_cost_stream_controls_change_run_identity() -> None:
    default = portfolio.load_economics_config()
    settings = {"start_year": default.start_year, "end_year": default.end_year,
                "real_discount_rate": default.real_discount_rate,
                "scghg_discount_rate": default.scghg_discount_rate,
                "demand_surge": default.demand_surge}
    full = portfolio._configured_run(default, settings,
                                     {CostStream.MATERIAL, CostStream.LABOR, CostStream.DISPOSAL, CostStream.CARBON})
    reduced = portfolio._configured_run(default, settings, {CostStream.MATERIAL, CostStream.LABOR})

    assert full.enabled_cost_streams != reduced.enabled_cost_streams
    assert portfolio._effective_identity("same-assets", full, settings) != portfolio._effective_identity(
        "same-assets", reduced, settings
    )
    assert CostStream.LOSS_OF_USE not in reduced.enabled_cost_streams


def test_filtered_workbook_preserves_other_sheets_and_source_columns(
    tmp_path: Path,
) -> None:
    portfolio = validate_portfolio(DEFAULT_WORKBOOK)
    selected = portfolio.display["asset_id"].head(2).to_list()
    destination = tmp_path / "selected.xlsx"

    materialize_filtered_workbook(DEFAULT_WORKBOOK, destination, selected)

    workbook = load_workbook(destination, read_only=True, data_only=True)
    assert "asset_validation_lists" in workbook.sheetnames
    assert [cell.value for cell in workbook["Sheet1"][1]] == portfolio.source.columns
    assert validate_portfolio(destination).source.height == 2


def test_analysis_identity_ignores_filter_order() -> None:
    checksum = sha256_file(DEFAULT_WORKBOOK)

    assert analysis_identity(checksum, ["A-2", "A-1"]) == analysis_identity(
        checksum, ["A-1", "A-2"]
    )


def test_invalid_year_settings_show_error_without_crashing(monkeypatch) -> None:
    from streamlit.testing.v1 import AppTest

    def portfolio_page():
        from cadence.ui.pages import portfolio

        portfolio.render()

    monkeypatch.setattr(portfolio, "_analysis_settings", lambda config: {
        "start_year": 2050, "end_year": 2026,
    })
    app = AppTest.from_function(portfolio_page)
    app.session_state["run_state"] = {"run_id": "stale"}
    app.run()
    assert not app.exception
    assert "End year must be no earlier than start year" in app.warning[0].value
    assert "run_state" not in app.session_state